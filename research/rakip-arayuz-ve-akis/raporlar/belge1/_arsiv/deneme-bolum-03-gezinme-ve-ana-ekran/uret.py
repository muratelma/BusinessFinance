from pathlib import Path
import re, shutil, hashlib, json
import pymupdf
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[2]
NAME = 'bolum-03'
ACCENT, INK, MUTED = '2F5D8A', '1F2A37', '5B6675'

# (kimlik, kanitlar/ altındaki özgün dosya, karartılacak alan: sol, üst, sağ, alt piksel)
ASSETS = [
    ('E0231', 'money-manager/07-rapor-agustos.png', None),
    ('E0236', 'money-manager/12-hesaplar-kart-borcu-bu-ay.png', None),
    ('E0254', 'money-manager/30-ayarlar-izgarasi.png', None),
    ('E0228', 'money-manager/03-dolu-ana-ekran.png', None),
    ('E0227', 'money-manager/02-bos-ana-ekran.png', None),
    ('E0017', 'bluecoins/01-ilk-acilis.png', None),
    ('E0018', 'bluecoins/02-bos-ana-ekran.png', None),
    ('E0025', 'bluecoins/09-ozgun-ozellik.png', None),
    ('E0026', 'bluecoins/10-fresh-bos-ana-ekran.png', None),
    ('E0056', 'bluecoins/f7-05-hatirlaticilar.png', None),
    ('E0084', 'bluecoins/f7-33-menu2.png', None),
    ('E0103', 'bluecoins/f7-52-nav-check.png', None),
    ('E0274', 'wallet-budgetbakers/00-magaza.png', None),
    ('E0275', 'wallet-budgetbakers/02b-menu.png', (200, 205, 560, 296)),
    ('E0276', 'wallet-budgetbakers/03-dolu-ana-ekran.png', None),
    ('E0376', 'wallet-budgetbakers/f7-55-menu-scroll.png', None),
    ('E0135', 'hesap-defterim/01-ilk-acilis-hosgeldin.png', None),
    ('E0136', 'hesap-defterim/02-bos-ana-ekran.png', None),
    ('E0137', 'hesap-defterim/03-dolu-ana-ekran.png', None),
    ('E0171', 'hesap-defterim/36-drawer-menu-ust.png', None),
    ('E0106', 'goodbudget/01-ilk-acilis.png', None),
    ('E0115', 'goodbudget/08-envelopes-filled-home.png', (160, 160, 450, 262)),
    ('E0211', 'kolaybi/d24-destek-guncel-durum-panosu.png', None),
    ('E0181', 'kolaybi/d33-video-guncel-durum-panosu.png', None),
    ('E0187', 'kolaybi/d39-guncel-arayuz-2026.png', None),
    ('E0262', 'parasut/04-video-cari-hesap-durumu.png', None),
]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


(BASE / 'kanit').mkdir(exist_ok=True)
rows = []
for ident, rel, box in ASSETS:
    src, dst = ROOT / 'kanitlar' / rel, BASE / 'kanit' / f'{ident}.png'
    if box:
        pix = pymupdf.Pixmap(str(src))
        if pix.alpha:
            pix = pymupdf.Pixmap(pix, 0)
        pix.set_rect(pymupdf.IRect(*box), (96, 96, 96))
        pix.save(str(dst))
    else:
        shutil.copyfile(src, dst)
        assert src.read_bytes() == dst.read_bytes()
    rows.append({'id': ident, 'original': rel, 'original_sha256': sha(src), 'copy_sha256': sha(dst),
                 'redacted_box': list(box) if box else None})

source = [
    '# Kaynak dizini — Belge 1 · Bölüm 3', '',
    '16 Eylül 2026. Görseller önceki Android koşumları ve resmî destek/video kaynaklarından; yeni koşum yapılmadı. Karartılan iki kopya dışındakiler özgün dosyayla birebir aynıdır.', '',
    '| Kimlik | Özgün dosya (araştırma kanitlar/ altında) | Kopya | Özgün SHA-256 | Not |', '|---|---|---|---|---|',
]
for r in rows:
    note = f"Kişisel ad karartıldı (piksel {r['redacted_box']}); kopya SHA-256 {r['copy_sha256']}" if r['redacted_box'] else 'Değiştirilmedi'
    source.append(f"| {r['id']} | {r['original']} | [Görsel](kanit/{r['id']}.png) | {r['original_sha256']} | {note} |")
(BASE / 'kaynaklar.md').write_text('\n'.join(source) + '\n', encoding='utf-8')
(BASE / 'kanit-manifest.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(1.6), Cm(1.5), Cm(1.8), Cm(1.8)
s.header_distance = s.footer_distance = Cm(.7)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(5), 1.05
for name, size, color, before in [('Title', 24, INK, 0), ('Heading 2', 15, ACCENT, 4), ('Heading 3', 12, INK, 8)]:
    st = doc.styles[name]
    st.font.name, st.font.size = 'Calibri', Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(4)
doc.styles['List Bullet'].font.size = Pt(10.5)
header = s.header.paragraphs[0]
header.text = 'BELGE 1 · RAKİP ARAYÜZ YAKLAŞIMLARI  ·  BÖLÜM 3'
header.style = 'Caption'
header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Taslak · 16 Eylül 2026   |   ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Belge 1 · Bölüm 3 · Gezinme ve ana ekran'
doc.core_properties.subject = 'Rakip arayüz yaklaşımları'
doc.core_properties.author = 'BusinessFinance'
doc.core_properties.keywords = 'Belge 1; bölüm 3; gezinme; ana ekran'


def inline(p, text, size=None, color=None):
    for part in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\))', text):
        if not part:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
        if link:
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), p.part.relate_to(link[2], RT.HYPERLINK, is_external=True))
            r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
            c = OxmlElement('w:color'); c.set(qn('w:val'), ACCENT); props.append(c)
            u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); props.append(u)
            r.append(props); t = OxmlElement('w:t'); t.text = link[1]; r.append(t); h.append(r); p._p.append(h)
            continue
        run = p.add_run(part[2:-2] if part.startswith('**') else part)
        run.bold = part.startswith('**') or None
        if size: run.font.size = size
        if color: run.font.color.rgb = RGBColor.from_string(color)


def no_split(table):
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))


def kucukler(lines):
    imgs = [re.fullmatch(r'!\[(.*?)\]\((.*?)\)', l.strip()) for l in lines]
    width = {5: 3.1, 4: 3.7, 3: 3.9, 2: 8.2}.get(len(imgs), 4.0)  # İkili şerit geniş masaüstü kareleri içindir.
    t = doc.add_table(rows=1, cols=len(imgs)); t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, img in zip(t.rows[0].cells, imgs):
        cell.width = Cm(17.4 / len(imgs))
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(BASE / img[2]), width=Cm(width))
        p.runs[-1]._r.xpath('.//wp:docPr')[0].set('descr', img[1])
        cp = cell.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after = Pt(4)
        inline(cp, img[1], Pt(8), MUTED)
    no_split(t)


def tablo(lines):
    vals = [[v.strip() for v in l.strip().strip('|').split('|')] for l in lines]
    vals = [v for v in vals if not all(re.fullmatch(r':?-+:?', x) for x in v)]
    size = Pt(9) if len(vals[0]) >= 4 else Pt(9.5)
    t = doc.add_table(rows=0, cols=len(vals[0])); t.style = 'Light List Accent 1'
    if len(vals[0]) == 3:
        t.autofit = False
        for col, width in zip(t.columns, (3.0, 8.0, 6.4)):
            col.width = Cm(width)
    for n, row in enumerate(vals):
        for cell, v in zip(t.add_row().cells, row):
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
            # Başlık ilk veri satırıyla tutulur; uzun tablolar satırlar arasında bölünebilir.
            p.paragraph_format.keep_with_next = n == 0
            inline(p, f'**{v}**' if n == 0 else v, size)
        if n == 0:
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    no_split(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


lines = (BASE / f'{NAME}.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line == '<!-- pagebreak -->':
        doc.add_page_break(); i += 1; continue
    if line == '<!-- kucukler -->':
        j = lines.index('<!-- /kucukler -->', i)
        kucukler([l for l in lines[i + 1:j] if l.strip()])
        i = j + 1; continue
    if line.startswith('|'):
        group = []
        while i < len(lines) and lines[i].startswith('|'):
            group.append(lines[i]); i += 1
        tablo(group); continue
    heading = re.match(r'^(#{1,3}) (.*)', line)
    if heading:
        level = len(heading[1])
        inline(doc.add_paragraph(style='Title' if level == 1 else f'Heading {level}'), heading[2])
    elif line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3); inline(p, line[2:])
    elif line.startswith('Şekil '):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; inline(p, line, Pt(8.5), MUTED)
    elif line.startswith('Belge 1 —'):
        inline(doc.add_paragraph(), line, Pt(9.5), MUTED)
    else:
        inline(doc.add_paragraph(), line)
    i += 1

out = BASE / f'{NAME}.docx'
doc.save(out)
print(f'DOCX üretildi: {out}; {len(ASSETS)} kanıt; kaynak dizini ve manifest yazıldı.')
