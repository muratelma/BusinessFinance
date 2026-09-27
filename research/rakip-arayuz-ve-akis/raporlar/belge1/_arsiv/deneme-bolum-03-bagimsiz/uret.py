from pathlib import Path
import re, json, hashlib, shutil
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import pymupdf

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[2]
REFERENCE=BASE.parent/'bolum-03-gezinme-ve-ana-ekran'
text=(BASE/'bolum-03.md').read_text(encoding='utf-8')
figures=re.findall(r'!\[([^]]+)\]\(([^)]+)\)',text)
selected={Path(path).stem for _,path in figures}
support={'E0231','E0236','E0254','E0376','E0181','E0018'}
old={x['id']:x for x in json.loads((REFERENCE/'kanit-manifest.json').read_text(encoding='utf-8'))}
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
(BASE/'kanit').mkdir(exist_ok=True)
manifest=[]
for ident in sorted(selected|support):
    prior=old[ident]
    original=ROOT/'kanitlar'/prior['original']
    assert sha(original)==prior['original_sha256']
    # Önceden karartılmış iki teslim kopyası aynen korunur; görüntü yeniden düzenlenmez.
    source=REFERENCE/'kanit'/f'{ident}.png' if prior['redacted_box'] else original
    assert sha(source)==prior['copy_sha256']
    target=BASE/'kanit'/f'{ident}.png'
    shutil.copyfile(source,target)
    manifest.append(dict(prior,role='Şekil' if ident in selected else 'Ek dayanak'))
(BASE/'kanit-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
source=['# Kaynak dizini','', '16 Eylül 2026. 18 basılı şekil; 6 ek dayanak. Özgün kanıtlar değiştirilmedi. E0115 ve E0275 mevcut karartılmış teslim kopyalarından alındı.','', '| Kimlik | Rol | Özgün kanıt | Teslim kopyası | Özgün SHA-256 |','|---|---|---|---|---|']
for x in manifest: source.append(f"| {x['id']} | {x['role']} | [Özgün](../../../kanitlar/{x['original']}) | [Kopya](kanit/{x['id']}.png) | {x['original_sha256']} |")
source+=['','Koşum kayıtları: [gözlemler](../../../gozlemler/README.md). Görsel envanteri: [KANIT-ENVANTERI.md](../../../KANIT-ENVANTERI.md). Karartma kutuları ve teslim kopyası hashleri manifesttedir.']
(BASE/'kaynaklar.md').write_text('\n'.join(source)+'\n',encoding='utf-8')

doc=Document();sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.left_margin=sec.right_margin=Cm(1.7)
sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.5)
sec.header_distance=sec.footer_distance=Cm(.65)
for name,size in [('Normal',10),('Title',26),('Heading 1',16),('Heading 2',13),('Caption',8.5)]:
    st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size)
    st.font.color.rgb=RGBColor.from_string('203244')
    st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.line_spacing=1.04
header=sec.header.paragraphs[0];header.text='RAKİP ARAYÜZ YAKLAŞIMLARI   /   03';header.style='Caption'
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
p.add_run('Bağımsız taslak · 16 Eylül 2026   /   ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
doc.core_properties.title='Bölüm 3 — Gezinme ve ana ekran — Bağımsız taslak'
doc.core_properties.author='BusinessFinance'

def inline(p,s,size=None):
    for part in re.split(r'(\*\*.*?\*\*)',s):
        run=p.add_run(part[2:-2] if part.startswith('**') else part)
        run.bold=part.startswith('**')
        if size:run.font.size=Pt(size)

def table(lines):
    rows=[[x.strip() for x in line.strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',v) for v in r)]
    t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Light Shading Accent 1';t.autofit=False
    widths=(3.0,7.5,7.1)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for n,values in enumerate(rows):
        row=t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,value in enumerate(values):
            row.cells[j].width=Cm(widths[j]);p=row.cells[j].paragraphs[0]
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=n==0
            inline(p,'**'+value+'**' if n==0 else value,8.7)
        if n==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def gallery(lines,wide=False):
    items=[re.fullmatch(r'!\[(.*?)\]\((.*?)\)',l) for l in lines if l.strip()]
    groups=[[x] for x in items] if wide else [items]
    for group in groups:
        t=doc.add_table(rows=1,cols=len(group));t.autofit=False
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,m in zip(t.rows[0].cells,group):
            cell.width=Cm(17.6/len(group));p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pix=pymupdf.Pixmap(str(BASE/m[2]));ratio=pix.width/pix.height
            h=min(6.7 if wide else (11.0 if len(group)==2 else 9.2),(17.0/len(group))/ratio)
            run=p.add_run();run.add_picture(str(BASE/m[2]),height=Cm(h))
            run._r.xpath('.//wp:docPr')[0].set('descr',m[1])
            cap=cell.add_paragraph();cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after=Pt(6);inline(cap,m[1],8.3)

lines=text.splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- page -->':doc.add_page_break();i+=1;continue
    if line in ('<!-- gallery -->','<!-- wide -->'):
        end='<!-- /wide -->' if 'wide' in line else '<!-- /gallery -->'
        j=lines.index(end,i);gallery(lines[i+1:j],'wide' in line);i=j+1;continue
    if line.startswith('|'):
        j=i
        while j<len(lines) and lines[j].startswith('|'):j+=1
        table(lines[i:j]);i=j;continue
    h=re.match(r'^(#{1,3}) (.*)',line)
    if h:
        p=doc.add_paragraph(style='Title' if len(h[1])==1 else 'Heading '+str(len(h[1])-1));inline(p,h[2])
    elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
    else:inline(doc.add_paragraph(),line)
    i+=1
doc.save(BASE/'bolum-03.docx')
assert len(doc.inline_shapes)==18
print('Word hazır: 18 şekil; 24 kanıt kopyası.')
